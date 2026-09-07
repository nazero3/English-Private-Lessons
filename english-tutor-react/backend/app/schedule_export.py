from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schedule_grid import WEEKDAYS, format_clock, occupied_starts


def _fill(hex_color: str) -> PatternFill:
    color = (hex_color or "#F5E6C8").lstrip("#")
    return PatternFill("solid", fgColor=color)


def build_teacher_workbook(teacher_name: str, slots: list[dict], time_starts: list[int]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Weekly"
    ws.sheet_view.rightToLeft = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:1"

    thin = Border(
        left=Side(style="thin", color="C4B48A"),
        right=Side(style="thin", color="C4B48A"),
        top=Side(style="thin", color="C4B48A"),
        bottom=Side(style="thin", color="C4B48A"),
    )
    header_fill = PatternFill("solid", fgColor="1A2656")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    time_font = Font(name="Calibri", bold=True, color="1A2656", size=11)
    cell_font = Font(name="Calibri", size=11, color="1A2656")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)

    headers = ["الوقت"] + [day["ar"] for day in WEEKDAYS]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    start_row = {minute: index + 2 for index, minute in enumerate(time_starts)}
    occupancy: dict[tuple[int, int], dict] = {}
    for slot in slots:
        for tick in occupied_starts(slot["start_minutes"], slot["duration_minutes"]):
            occupancy[(slot["weekday"], tick)] = slot

    skip: set[tuple[int, int]] = set()
    for minute in time_starts:
        row_idx = start_row[minute]
        time_cell = ws.cell(row_idx, 1, format_clock(minute))
        time_cell.font = time_font
        time_cell.alignment = center
        time_cell.fill = PatternFill("solid", fgColor="FFF8E8")
        time_cell.border = thin
        for day in WEEKDAYS:
            col_idx = day["id"] + 2
            key = (day["id"], minute)
            if key in skip:
                continue
            slot = occupancy.get(key)
            cell = ws.cell(row_idx, col_idx)
            cell.border = thin
            cell.alignment = center
            if not slot or slot["start_minutes"] != minute:
                continue
            span = slot["duration_minutes"] // 30
            label = slot.get("label") or f"{slot.get('student_name', '')} / {teacher_name}"
            cell.value = label
            cell.font = cell_font
            cell.fill = _fill(slot.get("color"))
            if span > 1:
                end_row = row_idx + span - 1
                ws.merge_cells(start_row=row_idx, start_column=col_idx, end_row=end_row, end_column=col_idx)
                for extra in range(1, span):
                    extra_min = minute + extra * 30
                    skip.add((day["id"], extra_min))
                    extra_cell = ws.cell(start_row[extra_min], col_idx)
                    extra_cell.border = thin
                    extra_cell.fill = _fill(slot.get("color"))

    ws.column_dimensions["A"].width = 12
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.row_dimensions[1].height = 28
    for minute in time_starts:
        ws.row_dimensions[start_row[minute]].height = 28

    ws.oddHeader.right.text = teacher_name
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
